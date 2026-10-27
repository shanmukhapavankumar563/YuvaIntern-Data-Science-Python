"""Mall customer K-Means clustering with reproducible evaluation and reporting.

Run from this directory: python clustering_analysis.py
The source file is cached in data/raw; generated artifacts are written to outputs/.
"""
from __future__ import annotations

import json
import urllib.request
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.metrics import silhouette_score
from sklearn.preprocessing import StandardScaler

HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"
RAW = HERE / "data" / "raw" / "Mall_Customers.csv"
SOURCE_URL = "https://raw.githubusercontent.com/MahimaSriK/Mall-Customer-Segmentation-Dataset/main/Mall_Customers.csv"
FEATURES = ["Age", "Annual Income (k$)", "Spending Score (1-100)"]
RANDOM_STATE = 42
MAX_K = 10
sns.set_theme(style="whitegrid", palette="deep")


def load_dataset() -> tuple[pd.DataFrame, str]:
    """Load cached public CSV or download it on the first run."""
    RAW.parent.mkdir(parents=True, exist_ok=True)
    if not RAW.exists():
        req = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "YuvaIntern-Week3-Clustering/1.0"})
        with urllib.request.urlopen(req, timeout=30) as response:
            RAW.write_bytes(response.read())
    return pd.read_csv(RAW), SOURCE_URL + " (cached at data/raw/Mall_Customers.csv)"


def find_column(df: pd.DataFrame, name: str) -> str:
    """Resolve common column-name variants without silently choosing unrelated fields."""
    normalized = {str(c).strip().lower().replace("_", " "): c for c in df.columns}
    aliases = {
        "CustomerID": ["customerid", "customer id"],
        "Gender": ["gender", "genre"],
        "Age": ["age"],
        "Annual Income (k$)": ["annual income (k$)", "annual income", "annual income (k $)", "annual income k$"],
        "Spending Score (1-100)": ["spending score (1-100)", "spending score", "spending score 1-100"],
    }
    keys = [x.lower().replace("_", " ") for x in aliases.get(name, [name])]
    for key in keys:
        if key in normalized:
            return normalized[key]
    raise ValueError(f"Required dataset field '{name}' not found. Available columns: {list(df.columns)}")


def add_table(doc: Document, frame: pd.DataFrame, max_rows: int = 20) -> None:
    frame = frame.head(max_rows)
    if frame.empty or len(frame.columns) == 0:
        doc.add_paragraph("No rows available.")
        return
    table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, frame.columns): cell.text = str(value)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = "—" if pd.isna(value) else (f"{value:.3f}" if isinstance(value, (float, np.floating)) else str(value))


def cluster_interpretations(profile: pd.DataFrame, overall_means: pd.Series) -> pd.DataFrame:
    """Describe each cluster relative to the sample mean, without assigning personas."""
    descriptions = []
    for cluster_id, row in profile.iterrows():
        statements = []
        for feature in FEATURES:
            value = row.get(f"{feature}_mean", np.nan)
            baseline = overall_means.get(feature, np.nan)
            if pd.isna(value) or pd.isna(baseline):
                continue
            direction = "above" if value > baseline else "below" if value < baseline else "at"
            statements.append(f"mean {feature} is {direction} the full-sample mean ({value:.2f} vs {baseline:.2f})")
        descriptions.append({"cluster": cluster_id, "interpretation": "; ".join(statements) + "."})
    return pd.DataFrame(descriptions)


def save_fig(fig: plt.Figure, name: str) -> Path:
    dest = OUT / name
    fig.tight_layout(); fig.savefig(dest, dpi=170, bbox_inches="tight"); plt.close(fig)
    return dest


def build_report(data: pd.DataFrame, source: str, quality: dict, stats: pd.DataFrame,
                 eval_table: pd.DataFrame, selected_k: int, selected_silhouette: float,
                 summary: pd.DataFrame, gender_profile: pd.DataFrame, interpretations: pd.DataFrame,
                 pca_variance: list[float],
                 charts: list[Path], feature_names: list[str], scaled: bool) -> None:
    doc = Document()
    doc.add_heading("Week 3: Unsupervised Learning and Clustering Analysis", 0)
    doc.add_paragraph("YuvaIntern Virtual Data Science with Python Trainee Internship")
    doc.add_heading("Introduction and objectives", 1)
    doc.add_paragraph("The objective is to explore a public customer dataset, prepare meaningful numeric features, compare K-Means cluster counts, and describe the resulting groups for exploratory segmentation. Clusters are statistical summaries, not verified customer types or causal explanations.")
    doc.add_heading("Dataset source and description", 1)
    doc.add_paragraph(f"Source: {source}. This Mall Customers dataset contains {quality['rows_after_dedup']} records and {quality['column_count']} fields. Its fields describe customer ID, gender, age, annual income (in thousands of dollars as labeled by the source), and a 1–100 spending score. The dataset is a small instructional sample; provenance and sampling details are limited.")
    doc.add_heading("Exploratory analysis and quality", 1)
    doc.add_paragraph(f"The original data had {quality['rows_before']} rows and {quality['column_count']} columns. There were {quality['duplicate_rows']} exact duplicate rows; {quality['duplicate_rows']} were removed. Missing cells before preprocessing: {quality['missing_cells_before']}. Numeric feature missing counts and column summaries are recorded in output CSVs. Negative age or income and spending scores outside 1–100 are counted and treated as invalid; missing and invalid numeric features, if any, were median-imputed. IQR flags are reported for review, not automatically removed.")
    doc.add_paragraph("Descriptive statistics:"); add_table(doc, stats.reset_index().rename(columns={"index": "feature"}))
    doc.add_paragraph("Data quality checks are saved in outputs/data_quality.json. The customer identifier is retained for traceability but excluded from clustering because its numeric values are labels, not behavioral measurements.")
    doc.add_heading("Feature selection and preprocessing", 1)
    doc.add_paragraph("Age, annual income, and spending score are used as clustering features because they represent different customer characteristics on interpretable numeric scales. CustomerID is excluded to prevent arbitrary ID distances from affecting clusters. Gender is kept for post-clustering profiling only; it is categorical and is not converted into an arbitrary numeric distance for this K-Means model. Missing feature values are median-imputed, exact full-row duplicates are removed, and StandardScaler standardizes each feature to mean zero and unit variance so the larger numeric scale does not dominate Euclidean distance.")
    doc.add_paragraph(f"Features actually used: {', '.join(feature_names)}. Scaling applied: {scaled}.")
    doc.add_heading("Key code excerpts", 2)
    doc.add_paragraph("The following excerpts show the preprocessing and model evaluation used by the reproducible script:")
    for snippet, explanation in [
        ("X_scaled = StandardScaler().fit_transform(X)\nmodel = KMeans(n_clusters=k, init='k-means++', n_init=20, random_state=42)\nlabels = model.fit_predict(X_scaled)", "Standardize the selected features, then fit deterministic K-Means candidates."),
        ("silhouette = silhouette_score(X_scaled, labels)\nselected_k = k_with_highest_valid_silhouette", "Compare candidate group separation with an internal score; the elbow plot remains a secondary, subjective check."),
    ]:
        doc.add_paragraph(explanation)
        paragraph = doc.add_paragraph(); run = paragraph.add_run(snippet); run.font.name = "Consolas"; run.font.size = Pt(9)
    doc.add_heading("Algorithm and cluster-count evaluation", 1)
    doc.add_paragraph("K-Means is a beginner-friendly centroid algorithm suitable for numeric, scaled features and compact groups. It minimizes within-cluster squared Euclidean distances. The fit uses k-means++ initialization, n_init=20, random_state=42, and max_iter=300. Candidate k values from 2 through 10 are compared using inertia (within-cluster sum of squared distances) and mean silhouette score. Inertia always decreases as k grows, so the elbow is subjective; silhouette prefers separated, compact groups and can favor simple partitions. Neither metric alone proves that clusters are meaningful.")
    doc.add_paragraph(f"For this run, the selected k is {selected_k}, chosen by the documented rule of the highest measured silhouette score among evaluated candidates ({selected_silhouette:.4f}). The elbow plot provides a secondary visual check; inertia decreases as k increases and the bend is subjective. The score and all candidate metrics are saved in outputs/cluster_evaluation.csv. Cluster labels are arbitrary identifiers.")
    add_table(doc, eval_table)
    doc.add_heading("Cluster results and profiles", 1)
    doc.add_paragraph("Profiles below summarize the original-unit feature values by assigned cluster. Gender proportions are included as descriptive context only; gender did not affect distance calculations.")
    add_table(doc, summary.reset_index())
    doc.add_heading("Cluster-by-cluster interpretation", 2)
    for _, row in interpretations.iterrows():
        doc.add_paragraph(f"Cluster {row['cluster']}: {row['interpretation']} These relative averages describe this sample and should not be treated as fixed customer personas.")
    if not gender_profile.empty:
        add_table(doc, gender_profile.reset_index())
    doc.add_paragraph("Possible applications are hypotheses for further research: a mall might test whether differentiated communications or service offers are useful for groups with different average income, age, or spending scores. These data do not include purchase histories, preferences, consent, or causal evidence, so no specific intervention or business outcome is established.")
    doc.add_heading("Visualizations", 1)
    descriptions = {
        "elbow_method.png": "Inertia versus k; a bend can suggest diminishing returns but does not uniquely identify the best k.",
        "silhouette_scores.png": "Mean silhouette score by k; higher scores indicate better separation under this metric, with caveats for non-spherical clusters and small samples.",
        "cluster_visualization.png": "Clusters shown in the income-spending plane. Age is used in fitting but not shown on this two-axis view, so points that overlap can still differ in age.",
        "cluster_pca_visualization.png": f"PCA projects standardized three-feature data into two dimensions for a compact view; the first two components explain {sum(pca_variance):.1%} of the variance. Projection can hide separation present in the full feature space.",
        "feature_distributions.png": "Feature distributions show spread and potential unusual values; unusual observations were retained unless they were exact duplicate rows.",
    }
    for chart in charts:
        doc.add_picture(str(chart), width=Inches(6.0)); doc.add_paragraph(descriptions.get(chart.name, chart.stem.replace("_", " ").title() + "."))
    doc.add_heading("Limitations and conclusion", 1)
    doc.add_paragraph("This dataset is small and its sampling method and score-generation process are not documented here. StandardScaler and K-Means favor roughly spherical clusters and are sensitive to outliers and chosen features. Silhouette and inertia are internal criteria, not external validation. The selected cluster count should be treated as an exploratory choice and checked against additional data and domain knowledge. Overall, this project demonstrates a reproducible clustering workflow and produces summaries that can guide questions for future research, not definitive customer segmentation decisions.")
    doc.save(OUT / "report.docx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    source_df, source = load_dataset()
    original_rows, original_cols = source_df.shape
    df = source_df.copy()
    # Resolve column aliases then rename to standard labels used consistently below.
    resolved = {standard: find_column(df, standard) for standard in FEATURES}
    for optional in ["CustomerID", "Gender"]:
        try:
            resolved[optional] = find_column(df, optional)
        except ValueError:
            # ID and gender are useful context but are not required model features.
            pass
    df = df.rename(columns={actual: standard for standard, actual in resolved.items()})
    missing_before = int(df.isna().sum().sum())
    missing_by_column = df.isna().sum().astype(int)
    duplicate_rows = int(df.duplicated().sum())
    df = df.drop_duplicates().copy()
    # Numeric conversion surfaces non-numeric strings as missing and logs them.
    conversion_invalid = {}
    for col in FEATURES:
        parsed = pd.to_numeric(df[col], errors="coerce")
        conversion_invalid[col] = int((df[col].notna() & parsed.isna()).sum())
        df[col] = parsed
    feature_missing = df[FEATURES].isna().sum()
    # Apply simple documented domain checks; unusual but possible observations remain.
    domain_invalid = {}
    for col, valid in [("Age", df["Age"].ge(0)),
                       ("Annual Income (k$)", df["Annual Income (k$)"].ge(0)),
                       ("Spending Score (1-100)", df["Spending Score (1-100)"].between(1, 100))]:
        invalid = df[col].notna() & ~valid
        domain_invalid[col] = int(invalid.sum())
        df.loc[invalid, col] = np.nan
    medians = df[FEATURES].median()
    # A wholly empty feature cannot be meaningfully imputed or clustered.
    empty_features = [col for col in FEATURES if pd.isna(medians[col])]
    if empty_features:
        raise ValueError(f"Cannot cluster: feature(s) contain no usable numeric values: {empty_features}")
    X = df[FEATURES].fillna(medians)
    outlier_rows = []
    for col in FEATURES:
        q1, q3 = X[col].quantile([.25, .75])
        iqr = q3 - q1
        lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        outlier_rows.append({"feature": col, "q1": float(q1), "q3": float(q3),
                             "lower_bound": float(lower), "upper_bound": float(upper),
                             "iqr_flagged_count": int(((X[col] < lower) | (X[col] > upper)).sum()),
                             "action": "Flag for review; no automatic removal"})
    pd.DataFrame(outlier_rows).to_csv(OUT / "outlier_review.csv", index=False)
    scaler = StandardScaler()
    X_scaled_array = scaler.fit_transform(X)
    X_scaled = pd.DataFrame(X_scaled_array, columns=FEATURES, index=df.index)
    quality = {"source": source, "rows_before": int(original_rows), "rows_after_dedup": int(len(df)),
               "column_count": int(original_cols), "columns": list(source_df.columns),
               "duplicate_rows": duplicate_rows, "missing_cells_before": missing_before,
               "missing_by_column_before": {str(k): int(v) for k, v in missing_by_column.items()},
               "feature_missing_before_imputation": {str(k): int(v) for k, v in feature_missing.items()},
               "non_numeric_feature_values_converted_to_missing": conversion_invalid,
               "domain_invalid_values_converted_to_missing": domain_invalid,
               "imputation_medians": {str(k): float(v) for k, v in medians.items()},
               "iqr_outlier_review": outlier_rows,
               "features_used": FEATURES, "gender_used_for_fit": False, "scaling": "StandardScaler"}
    (OUT / "data_quality.json").write_text(json.dumps(quality, indent=2), encoding="utf-8")
    source_df.describe(include="all").transpose().to_csv(OUT / "descriptive_statistics.csv")
    source_df.isna().sum().rename("missing_count").to_csv(OUT / "missing_values.csv")
    source_df.head(10).to_csv(OUT / "sample_records.csv", index=False)
    # Save auditable prepared features, both in original units and standardized units.
    prepared = df[[c for c in ["CustomerID", "Gender"] if c in df]].copy()
    for col in FEATURES: prepared[col] = X[col]
    prepared.to_csv(OUT / "prepared_customers.csv", index=False)
    X_scaled.add_prefix("scaled_").to_csv(OUT / "scaled_features.csv", index=False)

    max_k = min(MAX_K, len(df) - 1)
    if max_k < 2:
        raise ValueError("At least three usable rows are required to compare clustering candidates.")
    eval_rows = []
    models = {}
    # Compare k candidates with deterministic initialization and internal metrics.
    for k in range(2, max_k + 1):
        model = KMeans(n_clusters=k, init="k-means++", n_init=20, max_iter=300, random_state=RANDOM_STATE)
        labels = model.fit_predict(X_scaled)
        models[k] = model
        unique_labels = np.unique(labels)
        score = float(silhouette_score(X_scaled, labels)) if 2 <= len(unique_labels) < len(df) else float("nan")
        eval_rows.append({"k": k, "inertia": float(model.inertia_), "silhouette_score": score,
                          "actual_clusters": int(len(unique_labels))})
    evaluation = pd.DataFrame(eval_rows)
    evaluation.to_csv(OUT / "cluster_evaluation.csv", index=False)
    valid_evaluation = evaluation.dropna(subset=["silhouette_score"])
    if valid_evaluation.empty:
        raise ValueError("Silhouette scores could not be computed: the data produced fewer than two valid clusters for every candidate k.")
    selected_k = int(valid_evaluation.loc[valid_evaluation["silhouette_score"].idxmax(), "k"])
    model = models[selected_k]
    labels = model.labels_
    clustered = df.copy(); clustered["cluster"] = labels
    clustered.to_csv(OUT / "cluster_assignments.csv", index=False)
    centers = pd.DataFrame(scaler.inverse_transform(model.cluster_centers_), columns=FEATURES)
    centers.index.name = "cluster"
    centers.to_csv(OUT / "cluster_centers_original_units.csv")
    summary = clustered.groupby("cluster")[FEATURES].agg(["count", "mean", "median", "std"])
    # Flatten columns for easier spreadsheet use.
    summary.columns = [f"{feature}_{stat}" for feature, stat in summary.columns]
    gender_profile = pd.crosstab(clustered["cluster"], clustered["Gender"], normalize="index") if "Gender" in clustered else pd.DataFrame()
    if not gender_profile.empty: gender_profile = gender_profile.add_prefix("gender_share_")
    profile = summary.join(gender_profile)
    profile.to_csv(OUT / "cluster_summary.csv")
    interpretations = cluster_interpretations(profile, X.mean())
    interpretations.to_csv(OUT / "cluster_interpretations.csv", index=False)
    pca = PCA(n_components=2, random_state=RANDOM_STATE)
    projected = pca.fit_transform(X_scaled)
    pca_variance = [float(x) for x in pca.explained_variance_ratio_]
    pca_df = pd.DataFrame({"PC1": projected[:, 0], "PC2": projected[:, 1], "cluster": labels.astype(str)})
    pca_df.to_csv(OUT / "pca_coordinates.csv", index=False)

    charts: list[Path] = []
    fig, ax = plt.subplots(figsize=(8, 5)); ax.plot(evaluation.k, evaluation.inertia, marker="o", color="#2878a5")
    ax.axvline(selected_k, color="#d97925", linestyle="--", label=f"Selected k={selected_k}")
    ax.set(title="K-Means Elbow Method", xlabel="Number of clusters (k)", ylabel="Inertia (within-cluster sum of squares)", xticks=evaluation.k); ax.legend()
    charts.append(save_fig(fig, "elbow_method.png"))
    fig, ax = plt.subplots(figsize=(8, 5)); ax.plot(evaluation.k, evaluation.silhouette_score, marker="o", color="#468b65")
    ax.axvline(selected_k, color="#d97925", linestyle="--", label=f"Selected k={selected_k}")
    ax.set(title="Silhouette Score by Cluster Count", xlabel="Number of clusters (k)", ylabel="Mean silhouette score", xticks=evaluation.k); ax.legend()
    charts.append(save_fig(fig, "silhouette_scores.png"))
    fig, ax = plt.subplots(figsize=(9, 6)); sns.scatterplot(data=clustered, x="Annual Income (k$)", y="Spending Score (1-100)", hue="cluster", palette="tab10", s=65, alpha=.85, ax=ax)
    ax.set(title=f"Mall Customers by K-Means Cluster (k={selected_k})", xlabel="Annual income (thousands of dollars, source units)", ylabel="Spending score (1–100)"); ax.legend(title="Cluster")
    charts.append(save_fig(fig, "cluster_visualization.png"))
    fig, ax = plt.subplots(figsize=(8, 6)); sns.scatterplot(data=pca_df, x="PC1", y="PC2", hue="cluster", palette="tab10", s=60, alpha=.85, ax=ax)
    ax.set(title="K-Means Clusters Projected onto Two PCA Components", xlabel=f"PC1 ({pca_variance[0]:.1%} variance)", ylabel=f"PC2 ({pca_variance[1]:.1%} variance)"); ax.legend(title="Cluster")
    charts.append(save_fig(fig, "cluster_pca_visualization.png"))
    fig, axes = plt.subplots(1, 3, figsize=(14, 4))
    for ax, col in zip(axes, FEATURES): sns.histplot(df[col], kde=True, bins=18, ax=ax, color="#5187a3"); ax.set_title(col); ax.set_xlabel(col)
    charts.append(save_fig(fig, "feature_distributions.png"))
    (OUT / "run_summary.json").write_text(json.dumps({"selected_k": selected_k,
        "selected_silhouette_score": float(evaluation.loc[evaluation.k == selected_k, "silhouette_score"].iloc[0]),
        "random_state": RANDOM_STATE, "n_init": 20, "max_iter": 300,
        "pca_explained_variance_ratio": pca_variance,
        "cluster_sizes": {str(k): int(v) for k, v in clustered.cluster.value_counts().sort_index().items()}}, indent=2), encoding="utf-8")
    build_report(df, source, quality, df[FEATURES].describe().transpose(), evaluation, selected_k,
                 float(evaluation.loc[evaluation.k == selected_k, "silhouette_score"].iloc[0]),
                 profile, gender_profile, interpretations, pca_variance, charts, FEATURES, True)
    print(f"Loaded {len(df)} customers; selected k={selected_k} by highest silhouette score")
    print(f"Artifacts written to {OUT}")


if __name__ == "__main__":
    main()
