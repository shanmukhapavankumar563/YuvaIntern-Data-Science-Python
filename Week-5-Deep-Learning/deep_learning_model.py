"""Fashion-MNIST image classification using a compact PyTorch CNN.

Run from this directory: python deep_learning_model.py
The official test set is used once after validation-based training decisions.
"""
from __future__ import annotations

import json
import os
import random
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import sklearn
import torch
from sklearn.metrics import (accuracy_score, classification_report,
                             confusion_matrix, ConfusionMatrixDisplay,
                             f1_score, recall_score)
from sklearn.model_selection import train_test_split
from torch import nn
from torch.utils.data import DataLoader, Subset
from torchvision import datasets, transforms

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"
OUT = HERE / "outputs"
SOURCE_URL = "https://github.com/zalandoresearch/fashion-mnist"
LICENSE_URL = "https://github.com/zalandoresearch/fashion-mnist/blob/master/LICENSE"
CLASS_NAMES = ["T-shirt/top", "Trouser", "Pullover", "Dress", "Coat",
               "Sandal", "Shirt", "Sneaker", "Bag", "Ankle boot"]
SEED = 42
VALIDATION_FRACTION = 0.10
BATCH_SIZE = 128
MAX_EPOCHS = 5
PATIENCE = 2
LEARNING_RATE = 1e-3
WEIGHT_DECAY = 1e-4


class FashionCNN(nn.Module):
    """Two small convolutional blocks followed by a regularized classifier."""
    def __init__(self) -> None:
        super().__init__()
        self.features = nn.Sequential(
            nn.Conv2d(1, 32, kernel_size=3, padding=1), nn.ReLU(), nn.MaxPool2d(2),
            nn.Conv2d(32, 64, kernel_size=3, padding=1), nn.ReLU(), nn.MaxPool2d(2),
        )
        self.classifier = nn.Sequential(
            nn.Flatten(), nn.Linear(64 * 7 * 7, 128), nn.ReLU(),
            nn.Dropout(p=0.30), nn.Linear(128, 10),
        )

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return self.classifier(self.features(x))


def seed_everything(seed: int) -> None:
    random.seed(seed); np.random.seed(seed); torch.manual_seed(seed)
    if torch.cuda.is_available(): torch.cuda.manual_seed_all(seed)


def save_architecture_diagram(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 3.8)); ax.axis("off")
    boxes = [
        (0.02, "Input\n1×28×28", "#D6EAF8"),
        (0.18, "Conv 1→32, 3×3\nReLU + MaxPool\n32×14×14", "#D5F5E3"),
        (0.39, "Conv 32→64, 3×3\nReLU + MaxPool\n64×7×7", "#D5F5E3"),
        (0.60, "Flatten\n3136 values", "#FCF3CF"),
        (0.76, "Dense 128 + ReLU\nDropout 0.30", "#FADBD8"),
        (0.91, "10 class\nlogits", "#E8DAEF"),
    ]
    for x, label, color in boxes:
        ax.text(x, .52, label, transform=ax.transAxes, ha="center", va="center", fontsize=9,
                bbox={"boxstyle": "round,pad=0.6", "facecolor": color, "edgecolor": "#51606d"})
    for start, end in [(0.08,.13),(.26,.34),(.48,.55),(.67,.72),(.83,.87)]:
        ax.annotate("", xy=(end,.52), xytext=(start,.52), xycoords=ax.transAxes,
                    arrowprops={"arrowstyle":"->", "lw":1.5, "color":"#34495e"})
    ax.set_title("Fashion-MNIST convolutional neural network", fontsize=14, weight="bold")
    fig.tight_layout(); fig.savefig(path, dpi=170, bbox_inches="tight"); plt.close(fig)


def evaluate_model(model: nn.Module, loader: DataLoader, device: torch.device) -> tuple[np.ndarray, np.ndarray, np.ndarray, float]:
    model.eval(); all_targets=[]; all_predictions=[]; all_probabilities=[]; total_loss=0.0; count=0
    criterion = nn.CrossEntropyLoss()
    with torch.no_grad():
        for images, labels in loader:
            images, labels = images.to(device), labels.to(device)
            logits = model(images); loss = criterion(logits, labels)
            probabilities = torch.softmax(logits, dim=1)
            predictions = logits.argmax(dim=1)
            all_targets.extend(labels.cpu().numpy()); all_predictions.extend(predictions.cpu().numpy())
            all_probabilities.extend(probabilities.cpu().numpy())
            total_loss += loss.item() * len(labels); count += len(labels)
    return np.asarray(all_targets), np.asarray(all_predictions), np.asarray(all_probabilities), total_loss / max(count, 1)


def train_epoch(model: nn.Module, loader: DataLoader, optimizer, criterion, device) -> tuple[float, float]:
    model.train(); total_loss=0.0; correct=0; total=0
    for images, labels in loader:
        images, labels = images.to(device), labels.to(device)
        optimizer.zero_grad(set_to_none=True)
        logits = model(images); loss = criterion(logits, labels); loss.backward(); optimizer.step()
        total_loss += loss.item() * len(labels); correct += (logits.argmax(dim=1) == labels).sum().item(); total += len(labels)
    return total_loss / total, correct / total


def make_docx_fallback(path: Path, sections: list[tuple[str, str]], tables: list[tuple[str, pd.DataFrame]], images: list[tuple[Path, str]]) -> None:
    """Create a basic valid DOCX with standard-library OOXML when python-docx is absent."""
    W="http://schemas.openxmlformats.org/wordprocessingml/2006/main"; R="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    WP="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"; A="http://schemas.openxmlformats.org/drawingml/2006/main"; PIC="http://schemas.openxmlformats.org/drawingml/2006/picture"
    REL="http://schemas.openxmlformats.org/package/2006/relationships"; CT="http://schemas.openxmlformats.org/package/2006/content-types"
    for prefix, uri in [("w",W),("r",R),("wp",WP),("a",A),("pic",PIC),("rel",REL),("ct",CT)]: ET.register_namespace(prefix,uri)
    q=lambda ns, tag:f"{{{ns}}}{tag}"
    document=ET.Element(q(W,"document")); body=ET.SubElement(document,q(W,"body")); image_rels=[]
    def add_p(text, style="Normal"):
        p=ET.SubElement(body,q(W,"p")); ppr=ET.SubElement(p,q(W,"pPr")); ET.SubElement(ppr,q(W,"pStyle"),{q(W,"val"):style})
        run=ET.SubElement(p,q(W,"r")); ET.SubElement(run,q(W,"t")).text=str(text)
    def add_table(frame):
        tbl=ET.SubElement(body,q(W,"tbl")); tblpr=ET.SubElement(tbl,q(W,"tblPr")); borders=ET.SubElement(tblpr,q(W,"tblBorders"))
        for edge in ["top","left","bottom","right","insideH","insideV"]: ET.SubElement(borders,q(W,edge),{q(W,"val"):"single",q(W,"sz"):"4",q(W,"color"):"AAB7B8"})
        rows=[list(frame.columns)]+frame.astype(object).values.tolist()
        for ri,row in enumerate(rows):
            tr=ET.SubElement(tbl,q(W,"tr"))
            for value in row:
                tc=ET.SubElement(tr,q(W,"tc")); p=ET.SubElement(tc,q(W,"p")); run=ET.SubElement(p,q(W,"r")); ET.SubElement(run,q(W,"t")).text="—" if pd.isna(value) else str(value)
                if ri==0: ET.SubElement(ET.SubElement(run,q(W,"rPr")),q(W,"b"))
    def add_image(file, caption, idx):
        rid=f"rIdImg{idx}"; name=f"pic{idx}.png"; blob=file.read_bytes(); image_rels.append((rid,name,blob))
        width=int.from_bytes(blob[16:20],"big"); height=int.from_bytes(blob[20:24],"big"); cx=5_800_000; cy=int(cx*height/max(width,1))
        if cy>5_000_000: cy=5_000_000; cx=int(cy*width/max(height,1))
        p=ET.SubElement(body,q(W,"p")); run=ET.SubElement(p,q(W,"r")); drawing=ET.SubElement(run,q(W,"drawing")); inline=ET.SubElement(drawing,q(WP,"inline"))
        ET.SubElement(inline,q(WP,"extent"),{"cx":str(cx),"cy":str(cy)}); ET.SubElement(inline,q(WP,"docPr"),{"id":str(idx),"name":name,"descr":caption})
        frame=ET.SubElement(inline,q(WP,"cNvGraphicFramePr")); ET.SubElement(frame,q(A,"graphicFrameLocks"),{"noChangeAspect":"1"})
        graphic=ET.SubElement(inline,q(A,"graphic")); gd=ET.SubElement(graphic,q(A,"graphicData"),{"uri":PIC}); pic=ET.SubElement(gd,q(PIC,"pic")); nv=ET.SubElement(pic,q(PIC,"nvPicPr"))
        ET.SubElement(nv,q(PIC,"cNvPr"),{"id":str(idx),"name":name}); ET.SubElement(nv,q(PIC,"cNvPicPr")); fill=ET.SubElement(pic,q(PIC,"blipFill")); ET.SubElement(fill,q(A,"blip"),{q(R,"embed"):rid}); stretch=ET.SubElement(fill,q(A,"stretch")); ET.SubElement(stretch,q(A,"fillRect"))
        sp=ET.SubElement(pic,q(PIC,"spPr")); xfrm=ET.SubElement(sp,q(A,"xfrm")); ET.SubElement(xfrm,q(A,"off"),{"x":"0","y":"0"}); ET.SubElement(xfrm,q(A,"ext"),{"cx":str(cx),"cy":str(cy)}); geom=ET.SubElement(sp,q(A,"prstGeom"),{"prst":"rect"}); ET.SubElement(geom,q(A,"avLst")); add_p(caption)
    for title, text in sections:
        add_p(title,"Title" if title.startswith("Week 5") else "Heading1"); add_p(text)
    for title, frame in tables:
        add_p(title,"Heading1"); add_table(frame)
    for idx,(file,caption) in enumerate(images,1): add_image(file,caption,idx)
    sect=ET.SubElement(body,q(W,"sectPr")); ET.SubElement(sect,q(W,"pgSz"),{q(W,"w"):"12240",q(W,"h"):"15840"}); ET.SubElement(sect,q(W,"pgMar"),{q(W,"top"):"900",q(W,"right"):"900",q(W,"bottom"):"900",q(W,"left"):"900"})
    styles=ET.Element(q(W,"styles"))
    for sid,size,color in [("Normal","21","222222"),("Title","34","174A6E"),("Heading1","28","174A6E")]:
        st=ET.SubElement(styles,q(W,"style"),{q(W,"type"):"paragraph",q(W,"styleId"):sid}); ET.SubElement(st,q(W,"name"),{q(W,"val"):sid}); rp=ET.SubElement(st,q(W,"rPr")); ET.SubElement(rp,q(W,"rFonts"),{q(W,"ascii"):"Calibri",q(W,"hAnsi"):"Calibri"}); ET.SubElement(rp,q(W,"sz"),{q(W,"val"):size}); ET.SubElement(rp,q(W,"color"),{q(W,"val"):color})
    rootrels=ET.Element(q(REL,"Relationships")); ET.SubElement(rootrels,q(REL,"Relationship"),{"Id":"rId1","Type":"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument","Target":"word/document.xml"})
    docrels=ET.Element(q(REL,"Relationships")); ET.SubElement(docrels,q(REL,"Relationship"),{"Id":"rIdStyles","Type":"http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles","Target":"styles.xml"})
    for rid,name,_ in image_rels: ET.SubElement(docrels,q(REL,"Relationship"),{"Id":rid,"Type":"http://schemas.openxmlformats.org/officeDocument/2006/relationships/image","Target":f"media/{name}"})
    types=ET.Element(q(CT,"Types")); ET.SubElement(types,q(CT,"Default"),{"Extension":"rels","ContentType":"application/vnd.openxmlformats-package.relationships+xml"}); ET.SubElement(types,q(CT,"Default"),{"Extension":"xml","ContentType":"application/xml"}); ET.SubElement(types,q(CT,"Default"),{"Extension":"png","ContentType":"image/png"}); ET.SubElement(types,q(CT,"Override"),{"PartName":"/word/document.xml","ContentType":"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"}); ET.SubElement(types,q(CT,"Override"),{"PartName":"/word/styles.xml","ContentType":"application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"})
    with ZipFile(path,"w",ZIP_DEFLATED) as z:
        for name,root in [("[Content_Types].xml",types),("_rels/.rels",rootrels),("word/document.xml",document),("word/styles.xml",styles),("word/_rels/document.xml.rels",docrels)]: z.writestr(name,ET.tostring(root,encoding="utf-8",xml_declaration=True))
        for _,name,blob in image_rels: z.writestr(f"word/media/{name}",blob)


def create_report(data: dict, output_path: Path) -> None:
    """Create detailed report through python-docx or stdlib OOXML fallback."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError:
        make_docx_fallback(output_path, data["sections"], data["tables"], data["images"])
        return
    doc=Document()
    for title,text in data["sections"]:
        if title.startswith("Week 5"): doc.add_heading(title,0)
        else: doc.add_heading(title,1); doc.add_paragraph(text)
        if title.startswith("Week 5"): doc.add_paragraph(text)
    for title,frame in data["tables"]:
        doc.add_heading(title,1); t=doc.add_table(rows=1,cols=len(frame.columns)); t.style="Light Shading Accent 1"
        for c,val in zip(t.rows[0].cells,frame.columns): c.text=str(val)
        for _,row in frame.iterrows():
            cells=t.add_row().cells
            for c,val in zip(cells,row): c.text=str(val)
    for path,caption in data["images"]:
        doc.add_picture(str(path),width=Inches(6.0)); doc.add_paragraph(caption)
    doc.save(output_path)


def create_report_content(metadata: dict, history: pd.DataFrame, metrics: dict,
                          class_report: str, architecture_path: Path,
                          history_path: Path, confusion_path: Path,
                          sample_path: Path, class_counts: dict,
                          best_epoch: int, stopped_early: bool) -> dict:
    last=history.iloc[-1]
    if stopped_early:
        overfit_note=f"Validation loss reached its lowest value at epoch {best_epoch}; it failed to improve for {metadata['patience']} consecutive checks, so early stopping restored that checkpoint. The difference between training and validation curves should still be monitored."
    elif history.val_loss.is_monotonic_decreasing:
        overfit_note=f"Validation loss decreased at every observed epoch and reached {last.val_loss:.4f} at epoch {best_epoch}; training loss was {last.train_loss:.4f}. There is no validation-loss evidence of overfitting within these {len(history)} epochs, but the short CPU-limited run cannot rule out overfitting with longer training."
    else:
        overfit_note=f"Validation loss reached its lowest value at epoch {best_epoch}; the final recorded validation loss was {last.val_loss:.4f} versus training loss {last.train_loss:.4f}. Training reached the configured epoch cap; the diverging or fluctuating curves warrant monitoring and could motivate stronger regularization or earlier stopping."
    sections=[
        ("Week 5: Deep Learning Application in Data Science","Fashion-MNIST clothing image classification with PyTorch."),
        ("Objective and problem statement","Classify each 28×28 grayscale clothing image into one of ten garment categories. Inputs are normalized image pixels; targets are integer labels 0–9; output is a ten-class score vector and predicted class. Convolutional layers are appropriate because they learn local spatial patterns such as edges and textures while sharing weights across an image. A simpler linear or tabular model on flattened pixels may be faster but does not use image locality as naturally; the CNN needs more computation and can still generalize poorly beyond this benchmark. This is an educational image-recognition example, not a commercial product-recognition guarantee."),
        ("Dataset source, dimensions, and licensing",f"Fashion-MNIST by Zalando Research, source: {SOURCE_URL}. The dataset has 60,000 training and 10,000 official test images, each 28×28 grayscale, across ten classes. It is released under the MIT License; see {LICENSE_URL}. This run split the official training partition into {metadata['train_count']} training and {metadata['validation_count']} validation images; the official test set of {metadata['test_count']} remained held out."),
        ("Data preparation and preprocessing",f"Pixel images are represented as 28×28 grayscale arrays with integer values 0–255. torchvision ToTensor converts them to float tensors in [0,1] and adds the channel dimension, giving 1×28×28 inputs. A seeded stratified split creates validation indices from the official training portion. Per-class counts are balanced; observed split counts are recorded in outputs/dataset_inspection.json. No test image was used for fitting, early stopping, or model design decisions."),
        ("Neural network architecture", "Two convolutional blocks learn spatial patterns: Conv2d 1→32 (3×3, padding 1), ReLU, 2×2 max pooling; then Conv2d 32→64 (3×3, padding 1), ReLU, max pooling. The 64×7×7 map is flattened, passed through a 128-unit ReLU layer and Dropout(0.30), then a 10-logit output layer. CrossEntropyLoss combines multiclass log-softmax and negative log likelihood; Adam uses learning rate 0.001 and weight decay 0.0001. The architecture diagram is included below."),
        ("Training configuration and methodology",f"Device: {metadata['device']}; CPU thread limit: {metadata['cpu_threads']}. Batch size {metadata['batch_size']}; maximum {metadata['max_epochs']} epochs; early-stopping patience {metadata['patience']}; best validation-loss checkpoint restored at epoch {best_epoch}. Random seed {metadata['seed']}. The model uses Dropout to reduce co-adaptation and Adam weight decay as L2-style regularization. Early stopping avoids continuing when validation loss stops improving. A CPU-friendly CNN was selected because PyTorch is installed and CUDA is unavailable."),
        ("Actual training and validation results",f"Training completed {len(history)} epochs. Best validation loss was {history.val_loss.min():.4f}; validation accuracy at the best-loss checkpoint was {history.loc[history.epoch == best_epoch, 'val_accuracy'].iloc[0]:.4f}. The plotted curves show train and validation loss/accuracy for each completed epoch. {overfit_note}"),
        ("Held-out test evaluation",f"On the official held-out test set, accuracy was {metrics['accuracy']:.4f} and macro F1 was {metrics['macro_f1']:.4f}. The classification report below includes per-class precision, recall, F1, and support. The lowest class recall was {metadata.get('lowest_recall_detail','not available')}; the most frequent off-diagonal confusions were {metadata.get('frequent_confusions','not available')}. These patterns can guide data or architecture improvements, but do not establish performance on new images or real catalog photos."),
        ("Challenges, limitations, and improvements","Images are small grayscale crops, unlike real-world product photographs with varied lighting, backgrounds, angles, and multiple objects. A compact CNN limits CPU time but may underfit more complex patterns; five epochs and early stopping constrain runtime. Fashion-MNIST is balanced and relatively clean, so test performance may overstate deployment results. Potential improvements include data augmentation, learning-rate scheduling, repeated seeds, a larger GPU-trained model, and validation on a more representative product-photo dataset."),
        ("Practical applications and ethical considerations","Image classification can support exploratory catalog tagging or educational computer-vision pipelines. This dataset does not establish accuracy on commercial inventory, accessibility impacts, or fairness across users and products. Validate on representative images before any operational use."),
        ("Conclusion","A compact convolutional network was trained with an isolated validation set and evaluated on the untouched official test partition. The artifacts preserve the real history, predictions, and class-level evaluation for review; results should be interpreted within the dataset and hardware limitations."),
        ("Code excerpts and reproducibility","# Normalize pixels and add the channel dimension\ntransform = transforms.ToTensor()\n\n# Training-only optimization; validation controls checkpointing\nloss = criterion(model(images), labels)\nloss.backward()\noptimizer.step()"),
        ("Test classification report",class_report),
    ]
    tables=[("Class counts in official training/test data",pd.DataFrame({"class":CLASS_NAMES,"train_count":[class_counts['train'][i] for i in range(10)],"validation_count":[class_counts['validation'][i] for i in range(10)],"test_count":[class_counts['test'][i] for i in range(10)]})),
            ("Epoch training history",history.round(4))]
    images=[(architecture_path,"CNN layer diagram with tensor dimensions."),(history_path,"Actual training and validation loss and accuracy per epoch; the vertical line marks the best validation-loss epoch."),(confusion_path,"Actual test-set confusion matrix with true labels on rows and predicted labels on columns."),(sample_path,"Examples from the held-out test set with true and predicted class labels.")]
    return {"sections":sections,"tables":tables,"images":images}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    seed_everything(SEED)
    cpu_threads=min(6,os.cpu_count() or 1); torch.set_num_threads(cpu_threads)
    device=torch.device("cuda" if torch.cuda.is_available() else "cpu")
    transform=transforms.ToTensor()  # Converts uint8 [0,255] images to float [0,1].
    train_full=datasets.FashionMNIST(root=str(DATA_DIR),train=True,download=True,transform=transform)
    test_data=datasets.FashionMNIST(root=str(DATA_DIR),train=False,download=True,transform=transform)
    all_labels=np.asarray(train_full.targets)
    train_indices,val_indices=train_test_split(np.arange(len(train_full)),test_size=VALIDATION_FRACTION,
        random_state=SEED,stratify=all_labels)
    train_data=Subset(train_full,train_indices.tolist()); val_data=Subset(train_full,val_indices.tolist())
    generator=torch.Generator().manual_seed(SEED)
    pin=device.type=="cuda"
    train_loader=DataLoader(train_data,batch_size=BATCH_SIZE,shuffle=True,num_workers=0,pin_memory=pin,generator=generator)
    val_loader=DataLoader(val_data,batch_size=BATCH_SIZE,shuffle=False,num_workers=0,pin_memory=pin)
    test_loader=DataLoader(test_data,batch_size=BATCH_SIZE,shuffle=False,num_workers=0,pin_memory=pin)
    split_labels={"train":all_labels[train_indices],"validation":all_labels[val_indices],"test":np.asarray(test_data.targets)}
    class_counts={part:{i:int((labels==i).sum()) for i in range(10)} for part,labels in split_labels.items()}
    metadata={"source":SOURCE_URL,"license":"MIT","license_url":LICENSE_URL,"framework":"PyTorch","torch_version":torch.__version__,"torchvision_version":__import__("torchvision").__version__,"device":str(device),"cuda_available":torch.cuda.is_available(),"cpu_threads":cpu_threads,"image_shape":[1,28,28],"class_names":CLASS_NAMES,"train_count":len(train_indices),"validation_count":len(val_indices),"test_count":len(test_data),"batch_size":BATCH_SIZE,"max_epochs":MAX_EPOCHS,"patience":PATIENCE,"learning_rate":LEARNING_RATE,"weight_decay":WEIGHT_DECAY,"seed":SEED,"class_counts":class_counts}
    (OUT/"dataset_inspection.json").write_text(json.dumps(metadata,indent=2),encoding="utf-8")
    model=FashionCNN().to(device)
    architecture_path=OUT/"architecture.png"; save_architecture_diagram(architecture_path)
    criterion=nn.CrossEntropyLoss(); optimizer=torch.optim.Adam(model.parameters(),lr=LEARNING_RATE,weight_decay=WEIGHT_DECAY)
    history_rows=[]; best_val=float("inf"); best_epoch=0; wait=0; checkpoint_path=OUT/"best_model.pt"; stopped_early=False
    for epoch in range(1,MAX_EPOCHS+1):
        train_loss,train_accuracy=train_epoch(model,train_loader,optimizer,criterion,device)
        yv,pv,_,val_loss=evaluate_model(model,val_loader,device); val_accuracy=accuracy_score(yv,pv)
        history_rows.append({"epoch":epoch,"train_loss":train_loss,"train_accuracy":train_accuracy,"val_loss":val_loss,"val_accuracy":float(val_accuracy)})
        print(f"Epoch {epoch}/{MAX_EPOCHS} - train loss {train_loss:.4f}, acc {train_accuracy:.4f}; val loss {val_loss:.4f}, acc {val_accuracy:.4f}",flush=True)
        if val_loss < best_val:
            best_val=val_loss; best_epoch=epoch; wait=0
            torch.save({"model_state_dict":model.state_dict(),"class_names":CLASS_NAMES,"seed":SEED,"best_epoch":best_epoch},checkpoint_path)
        else:
            wait+=1
            if wait>=PATIENCE:
                stopped_early=True; print("Early stopping: validation loss did not improve.",flush=True); break
    history=pd.DataFrame(history_rows); history.to_csv(OUT/"training_history.csv",index=False)
    history_path=OUT/"training_history.png"
    fig,axes=plt.subplots(1,2,figsize=(12,4.5))
    axes[0].plot(history.epoch,history.train_loss,marker="o",label="Train"); axes[0].plot(history.epoch,history.val_loss,marker="o",label="Validation"); axes[0].axvline(best_epoch,color="#777",linestyle="--",label=f"Best epoch {best_epoch}"); axes[0].set(title="Cross-entropy loss",xlabel="Epoch",ylabel="Loss"); axes[0].legend()
    axes[1].plot(history.epoch,history.train_accuracy,marker="o",label="Train"); axes[1].plot(history.epoch,history.val_accuracy,marker="o",label="Validation"); axes[1].axvline(best_epoch,color="#777",linestyle="--",label=f"Best epoch {best_epoch}"); axes[1].set(title="Classification accuracy",xlabel="Epoch",ylabel="Accuracy",ylim=(0,1)); axes[1].legend()
    fig.suptitle("Fashion-MNIST training history"); fig.tight_layout(); fig.savefig(history_path,dpi=170,bbox_inches="tight"); plt.close(fig)
    best_checkpoint=torch.load(checkpoint_path,map_location=device,weights_only=True); model.load_state_dict(best_checkpoint["model_state_dict"])
    yt,yp,probabilities,test_loss=evaluate_model(model,test_loader,device)
    metrics={"test_loss":float(test_loss),"accuracy":float(accuracy_score(yt,yp)),"macro_f1":float(f1_score(yt,yp,average="macro")),"best_epoch":best_epoch,"epochs_completed":len(history),"early_stopped":stopped_early}
    class_recalls=recall_score(yt,yp,labels=list(range(10)),average=None,zero_division=0)
    weakest=int(np.argmin(class_recalls)); metrics["per_class_recall"]={CLASS_NAMES[i]:float(class_recalls[i]) for i in range(10)}
    metadata["lowest_recall_detail"]=f"{CLASS_NAMES[weakest]} ({class_recalls[weakest]:.2f})"
    report=classification_report(yt,yp,labels=list(range(10)),target_names=CLASS_NAMES,zero_division=0)
    (OUT/"classification_report.txt").write_text(report,encoding="utf-8")
    (OUT/"test_metrics.json").write_text(json.dumps(metrics,indent=2),encoding="utf-8")
    matrix=confusion_matrix(yt,yp,labels=list(range(10))); pd.DataFrame(matrix,index=CLASS_NAMES,columns=CLASS_NAMES).to_csv(OUT/"confusion_matrix.csv")
    confusion_without_diagonal=matrix.copy(); np.fill_diagonal(confusion_without_diagonal,0)
    frequent_errors=[]
    for _ in range(2):
        i,j=np.unravel_index(np.argmax(confusion_without_diagonal),confusion_without_diagonal.shape)
        frequent_errors.append(f"{CLASS_NAMES[i]} predicted as {CLASS_NAMES[j]} ({int(confusion_without_diagonal[i,j])})")
        confusion_without_diagonal[i,j]=0
    metadata["frequent_confusions"]="; ".join(frequent_errors)
    fig,ax=plt.subplots(figsize=(9,8)); disp=ConfusionMatrixDisplay(matrix,display_labels=CLASS_NAMES); disp.plot(ax=ax,xticks_rotation=45,cmap="Blues",colorbar=False,values_format="d"); ax.set_title("Fashion-MNIST held-out test confusion matrix"); fig.tight_layout(); confusion_path=OUT/"confusion_matrix.png"; fig.savefig(confusion_path,dpi=170,bbox_inches="tight"); plt.close(fig)
    pd.DataFrame({"image_index":np.arange(len(yt)),"true_label":yt,"true_class":[CLASS_NAMES[i] for i in yt],"predicted_label":yp,"predicted_class":[CLASS_NAMES[i] for i in yp],"confidence":probabilities.max(axis=1)}).to_csv(OUT/"test_predictions.csv",index=False)
    sample_path=OUT/"test_prediction_examples.png"; sample_indices=np.linspace(0,len(test_data)-1,10,dtype=int); fig,axes=plt.subplots(2,5,figsize=(12,5))
    for ax,i in zip(axes.flat,sample_indices):
        image,label=test_data[i]; ax.imshow(image.squeeze(0),cmap="gray"); ax.set_title(f"True: {CLASS_NAMES[yt[i]]}\nPred: {CLASS_NAMES[yp[i]]}",fontsize=8,color=("green" if yt[i]==yp[i] else "red")); ax.axis("off")
    fig.suptitle("Held-out Fashion-MNIST examples"); fig.tight_layout(); fig.savefig(sample_path,dpi=170,bbox_inches="tight"); plt.close(fig)
    best_checkpoint["test_accuracy"]=metrics["accuracy"]; best_checkpoint["test_macro_f1"]=metrics["macro_f1"]; torch.save(best_checkpoint,OUT/"fashion_mnist_cnn.pt")
    sections_data={"train":class_counts["train"],"validation":class_counts["validation"],"test":class_counts["test"]}
    metadata["class_counts"]=sections_data
    doc_data=create_report_content(metadata,history,metrics,report,architecture_path,history_path,confusion_path,sample_path,class_counts,best_epoch,stopped_early)
    create_report(doc_data,OUT/"report.docx")
    print(f"Test accuracy: {metrics['accuracy']:.4f}; macro F1: {metrics['macro_f1']:.4f}",flush=True)
    print(f"Best validation-loss epoch: {best_epoch}; outputs: {OUT}",flush=True)


def regenerate_final_artifacts() -> None:
    """Finish report/model packaging from artifacts after training has completed."""
    metadata=json.loads((OUT/"dataset_inspection.json").read_text(encoding="utf-8"))
    metrics=json.loads((OUT/"test_metrics.json").read_text(encoding="utf-8"))
    history=pd.read_csv(OUT/"training_history.csv")
    class_report=(OUT/"classification_report.txt").read_text(encoding="utf-8")
    predictions=pd.read_csv(OUT/"test_predictions.csv")
    true_labels=predictions["true_label"].to_numpy(); predicted_labels=predictions["predicted_label"].to_numpy()
    class_recalls=recall_score(true_labels,predicted_labels,labels=list(range(10)),average=None,zero_division=0)
    metrics["per_class_recall"]={CLASS_NAMES[i]:float(class_recalls[i]) for i in range(10)}
    weakest=int(np.argmin(class_recalls)); metadata["lowest_recall_detail"]=f"{CLASS_NAMES[weakest]} ({class_recalls[weakest]:.2f})"
    matrix=pd.read_csv(OUT/"confusion_matrix.csv",index_col=0).to_numpy(); np.fill_diagonal(matrix,0)
    frequent=[]
    for _ in range(2):
        i,j=np.unravel_index(np.argmax(matrix),matrix.shape); frequent.append(f"{CLASS_NAMES[i]} predicted as {CLASS_NAMES[j]} ({int(matrix[i,j])})"); matrix[i,j]=0
    metadata["frequent_confusions"]="; ".join(frequent)
    (OUT/"test_metrics.json").write_text(json.dumps(metrics,indent=2),encoding="utf-8")
    best_checkpoint=torch.load(OUT/"best_model.pt",map_location="cpu",weights_only=True)
    best_checkpoint["test_accuracy"]=metrics["accuracy"]; best_checkpoint["test_macro_f1"]=metrics["macro_f1"]
    torch.save(best_checkpoint,OUT/"fashion_mnist_cnn.pt")
    class_counts={part:{int(k):int(v) for k,v in counts.items()} for part,counts in metadata["class_counts"].items()}
    content=create_report_content(metadata,history,metrics,class_report,OUT/"architecture.png",
        OUT/"training_history.png",OUT/"confusion_matrix.png",OUT/"test_prediction_examples.png",
        class_counts,int(metrics["best_epoch"]),bool(metrics["early_stopped"]))
    create_report(content,OUT/"report.docx")
    print(f"Final model checkpoint and DOCX report regenerated from completed run: {OUT}")


if __name__=="__main__":
    regenerate_final_artifacts() if "--report-only" in sys.argv else main()
