"""Reproducible Titanic data acquisition, cleaning, visualisation and reporting.

Run from this folder with: python titanic_data_cleaning.py
The script keeps the source CSV and writes derived artifacts under outputs/.
"""
from __future__ import annotations

import json
import urllib.request
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # Allows chart creation in terminals and CI environments.
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
RAW_DIR = ROOT / "data" / "raw"
OUT = ROOT / "outputs"
SOURCE_URL = "https://raw.githubusercontent.com/mwaskom/seaborn-data/master/titanic.csv"
RAW_FILE = RAW_DIR / "titanic.csv"


def acquire_data() -> pd.DataFrame:
    """Download the openly available seaborn-data CSV once, then use local copy."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    if not RAW_FILE.exists():
        request = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "Week1-Data-Cleaning-Project/1.0"})
        with urllib.request.urlopen(request, timeout=30) as response:
            RAW_FILE.write_bytes(response.read())
    return pd.read_csv(RAW_FILE)


def quality_summary(df: pd.DataFrame) -> dict:
    return {
        "rows": int(df.shape[0]), "columns": int(df.shape[1]),
        "duplicate_rows": int(df.duplicated().sum()),
        "missing_cells": int(df.isna().sum().sum()),
        "missing_by_column": {str(k): int(v) for k, v in df.isna().sum().items()},
        "dtypes": {str(k): str(v) for k, v in df.dtypes.items()},
    }


def iqr_outliers(df: pd.DataFrame) -> dict:
    """Count Tukey IQR flags for numeric fields; flags are investigated, not removed."""
    result = {}
    for col in df.select_dtypes(include="number").columns:
        values = df[col].dropna()
        q1, q3 = values.quantile([0.25, 0.75])
        low, high = q1 - 1.5 * (q3 - q1), q3 + 1.5 * (q3 - q1)
        mask = values.lt(low) | values.gt(high)
        result[col] = {"q1": float(q1), "q3": float(q3), "lower_bound": float(low),
                       "upper_bound": float(high), "flagged_count": int(mask.sum()),
                       "flagged_row_indices": [int(i) for i in values.index[mask][:100]]}
    return result


def make_charts(before: pd.DataFrame, after: pd.DataFrame) -> list[Path]:
    charts = []
    missing = pd.DataFrame({"Before": before.isna().sum(), "After": after.isna().sum()})
    missing = missing.loc[missing.max(axis=1) > 0]
    if not missing.empty:
        ax = missing.plot(kind="bar", figsize=(11, 5), color=["#e67e22", "#2874a6"])
        ax.set_title("Missing values by column: before and after cleaning")
        ax.set_ylabel("Missing cells"); ax.set_xlabel("Column"); plt.xticks(rotation=35, ha="right")
        plt.tight_layout(); path = OUT / "missing_values_before_after.png"; plt.savefig(path, dpi=160); plt.close(); charts.append(path)
    numeric = [c for c in ["age", "fare"] if c in before]
    if numeric:
        fig, axes = plt.subplots(1, len(numeric), figsize=(6 * len(numeric), 4))
        if len(numeric) == 1: axes = [axes]
        for ax, col in zip(axes, numeric):
            ax.boxplot(before[col].dropna(), patch_artist=True, boxprops={"facecolor": "#85c1e9"})
            ax.set_title(f"{col.title()} distribution (source data)"); ax.set_ylabel(col.title())
        fig.tight_layout(); path = OUT / "numeric_outlier_review.png"; fig.savefig(path, dpi=160); plt.close(fig); charts.append(path)
    if "survived" in after:
        ax = after["survived"].value_counts().sort_index().plot(kind="bar", color=["#c0392b", "#27ae60"])
        ax.set_title("Passenger survival counts (cleaned dataset)"); ax.set_xlabel("Survived (0 = no, 1 = yes)"); ax.set_ylabel("Passengers")
        plt.tight_layout(); path = OUT / "survival_counts.png"; plt.savefig(path, dpi=160); plt.close(); charts.append(path)
    return charts


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers): cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row): cell.text = str(value)


def create_report(raw: pd.DataFrame, clean: pd.DataFrame, before: dict, after: dict,
                  outliers: dict, charts: list[Path]) -> None:
    doc = Document()
    doc.add_heading("Week 1: Data Acquisition, Cleaning, and Preprocessing", 0)
    doc.add_paragraph("YuvaIntern Virtual Data Science with Python Trainee Internship")
    doc.add_heading("Objective", 1)
    doc.add_paragraph("Acquire a public dataset, inspect its structure and quality, document defensible cleaning choices, and produce a reproducible cleaned dataset with supporting evidence.")
    doc.add_heading("Dataset source and acquisition", 1)
    doc.add_paragraph("Titanic passenger data from the seaborn-data public GitHub repository. The script downloads the CSV over HTTPS on its first run and retains it at data/raw/titanic.csv; subsequent runs use that local source copy. This source is a commonly used teaching dataset and includes passenger attributes and outcome fields.")
    doc.add_paragraph(SOURCE_URL)
    doc.add_heading("Initial exploration", 1)
    doc.add_paragraph(f"The source file has {before['rows']} rows and {before['columns']} columns. Fields: {', '.join(raw.columns)}.")
    doc.add_paragraph("The full data types and descriptive statistics are saved in outputs/data_types.csv and outputs/descriptive_statistics.csv. Missing-value counts are in outputs/missing_values.csv.")
    add_table(doc, ["Check", "Before", "After"], [["Rows", str(before["rows"]), str(after["rows"])], ["Columns", str(before["columns"]), str(after["columns"])], ["Duplicate rows", str(before["duplicate_rows"]), str(after["duplicate_rows"])], ["Missing cells", str(before["missing_cells"]), str(after["missing_cells"])]] )
    doc.add_heading("Data-quality issues and decisions", 1)
    doc.add_paragraph(f"Exact duplicate rows: {before['duplicate_rows']} were found. Exact repeats were removed because they duplicate complete records; no rows were removed for partial similarity.")
    doc.add_paragraph("Missing numeric age values are filled with the observed median, which is less sensitive than the mean to the skewed age distribution. Missing embarked values use the most frequent observed category. Missing deck values are replaced with the explicit label 'Unknown' because missingness may itself reflect unrecorded cabin information. embark_town is derived from embarked where possible and otherwise set to 'Unknown'. Other categorical gaps use the explicit 'Unknown' label, and any other numeric gaps use that column's median. Original columns are retained.")
    doc.add_paragraph("The code checks negative ages/fares and invalid values in the binary survived field. Such entries, if present, are counted and set to missing before the field-specific imputation; valid observations are retained. No unverified domain limits are imposed on age or fare.")
    doc.add_paragraph("Numeric fields are screened with Tukey's 1.5×IQR rule. This is a flagging method, not proof of error: high fares or ages can be valid passenger observations. No IQR-flagged observations are automatically deleted or capped. Exact counts and thresholds are written to outputs/iqr_outlier_review.json.")
    doc.add_paragraph("Before/after missing counts and duplicate counts are reported above; row-level categorical inconsistencies are checked for surrounding whitespace and normalized. The script records issue counts and actions in outputs/cleaning_log.json.")
    doc.add_heading("Method and reproducibility", 1)
    doc.add_paragraph("Run python titanic_data_cleaning.py from this folder. The script uses pandas for inspection and preprocessing, matplotlib for figures, and python-docx for this report. Source acquisition is HTTPS-based and the downloaded input is kept separately from derived files.")
    for snippet, explanation in [
        ("duplicate_count = data.duplicated().sum()\ndata = data.drop_duplicates()", "Count exact complete-row repeats before removing them."),
        ("data['age'] = data['age'].fillna(data['age'].median())", "Use a robust central value for missing numeric age."),
        ("lower = q1 - 1.5 * iqr\nupper = q3 + 1.5 * iqr", "Flag unusually distant numeric values for review, without automatic deletion."),
    ]:
        doc.add_paragraph(explanation)
        p = doc.add_paragraph(); run = p.add_run(snippet); run.font.name = "Consolas"; run.font.size = Pt(9)
    doc.add_heading("Before-and-after findings and charts", 1)
    doc.add_paragraph(f"Cleaning changed the row count from {before['rows']} to {after['rows']}; missing cells changed from {before['missing_cells']} to {after['missing_cells']}. Review the accompanying CSV and JSON summaries for column-level detail.")
    for chart in charts:
        doc.add_picture(str(chart), width=Inches(6.2)); doc.add_paragraph(chart.stem.replace("_", " ").title())
    doc.add_heading("Challenges and solutions", 1)
    doc.add_paragraph("Passenger records contain naturally missing values and skewed numeric fields. The workflow preserves useful columns, makes imputation choices explicit, stores the raw file unchanged, and treats statistical outliers as review candidates rather than errors.")
    doc.add_heading("Potential effects on future analysis", 1)
    doc.add_paragraph("Median imputation can reduce variance and introduce a concentration at the median; an Unknown category preserves missingness information but may become a predictive category. Removing exact duplicates prevents repeated identical rows from receiving extra weight. Retaining valid outliers preserves real variation, though some models may be sensitive to extreme values. For predictive modeling, imputation should be fitted on training data only to prevent leakage.")
    doc.add_heading("Conclusion", 1)
    doc.add_paragraph("This pipeline creates an auditable and reproducible cleaned Titanic dataset while retaining meaningful observations. Generated totals in this report are calculated directly from the local input at run time.")
    doc.save(OUT / "report.docx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    raw = acquire_data()
    before = quality_summary(raw)
    raw.describe(include="all").transpose().to_csv(OUT / "descriptive_statistics.csv")
    pd.DataFrame({"column": raw.columns, "dtype": [str(x) for x in raw.dtypes], "non_null": raw.notna().sum().values, "missing": raw.isna().sum().values}).to_csv(OUT / "data_types.csv", index=False)
    raw.isna().sum().rename("missing_count").to_csv(OUT / "missing_values.csv")
    outliers = iqr_outliers(raw)
    (OUT / "iqr_outlier_review.json").write_text(json.dumps(outliers, indent=2), encoding="utf-8")

    clean = raw.copy()
    duplicate_count = int(clean.duplicated().sum())
    clean = clean.drop_duplicates().copy()
    actions = {"exact_duplicate_rows_removed": duplicate_count, "invalid_values": {}}
    # Normalize whitespace only in text fields; preserve original category spellings otherwise.
    for col in clean.select_dtypes(include="object").columns:
        clean[col] = clean[col].map(lambda x: x.strip() if isinstance(x, str) else x)
    if "age" in clean:
        invalid = clean["age"].lt(0).fillna(False); actions["invalid_values"]["age_negative"] = int(invalid.sum()); clean.loc[invalid, "age"] = pd.NA
    if "fare" in clean:
        invalid = clean["fare"].lt(0).fillna(False); actions["invalid_values"]["fare_negative"] = int(invalid.sum()); clean.loc[invalid, "fare"] = pd.NA
    if "survived" in clean:
        invalid = ~clean["survived"].isin([0, 1]) & clean["survived"].notna(); actions["invalid_values"]["survived_not_0_or_1"] = int(invalid.sum()); clean.loc[invalid, "survived"] = pd.NA
    if "age" in clean: clean["age"] = clean["age"].fillna(clean["age"].median())
    if "fare" in clean: clean["fare"] = clean["fare"].fillna(clean["fare"].median())
    if "embarked" in clean:
        mode = clean["embarked"].mode(dropna=True)
        if not mode.empty: clean["embarked"] = clean["embarked"].fillna(mode.iloc[0])
    if "embark_town" in clean and "embarked" in clean:
        mapping = {"C": "Cherbourg", "Q": "Queenstown", "S": "Southampton"}
        clean["embark_town"] = clean["embark_town"].fillna(clean["embarked"].map(mapping))
    for col in clean.select_dtypes(include="object").columns:
        clean[col] = clean[col].fillna("Unknown")
    for col in clean.select_dtypes(include="number").columns:
        if clean[col].isna().any(): clean[col] = clean[col].fillna(clean[col].median())
    after = quality_summary(clean)
    clean.to_csv(OUT / "titanic_cleaned.csv", index=False)
    (OUT / "quality_summary_before_after.json").write_text(json.dumps({"before": before, "after": after}, indent=2), encoding="utf-8")
    (OUT / "cleaning_log.json").write_text(json.dumps({"source_url": SOURCE_URL, "actions": actions,
        "text_whitespace_normalization": "Applied to object columns; category spellings otherwise retained.",
        "outlier_policy": "IQR flags are reported; no outlier rows were automatically deleted or capped."}, indent=2), encoding="utf-8")
    charts = make_charts(raw, clean)
    create_report(raw, clean, before, after, outliers, charts)
    print(f"Input: {raw.shape[0]} rows x {raw.shape[1]} columns")
    print(f"Output: {clean.shape[0]} rows x {clean.shape[1]} columns; {after['missing_cells']} missing cells")
    print(f"Artifacts written to: {OUT}")


if __name__ == "__main__":
    main()
